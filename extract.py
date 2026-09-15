import re
import sys
import traceback
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tkinter import (
    Tk,
    Toplevel,
    StringVar,
    BooleanVar,
    filedialog,
    messagebox,
    Label,
    Button
)
from tkinter import ttk

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side
)


# ============================================================
# ОБЩИЕ ФУНКЦИИ
# ============================================================

def application_directory() -> Path:
    """Папка, где находится EXE или extract.py."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent

    return Path(__file__).resolve().parent


def clean_spaces(value: str) -> str:
    """Убирает повторяющиеся пробелы и переносы."""
    if not value:
        return ""

    value = value.replace("\u00a0", " ")
    value = value.replace("\r", " ")
    value = value.replace("\n", " ")
    value = re.sub(r"\s+", " ", value)

    return value.strip()


def normalize_text(value: str) -> str:
    """Нормализует текст для регулярных выражений."""
    if not value:
        return ""

    value = value.replace("\u00a0", " ")
    value = value.replace("–", "-")
    value = value.replace("—", "-")
    value = value.replace("−", "-")
    value = re.sub(r"[ \t]+", " ", value)

    return value


def digits_only(value: str) -> str:
    """Оставляет только цифры."""
    return re.sub(r"\D", "", value or "")


def normalize_money(value: str) -> str:
    """Приводит сумму к формату 12345.67."""
    value = value.strip()
    value = value.replace("\u00a0", "")
    value = value.replace(" ", "")
    value = value.replace(",", ".")

    return value


def parse_decimal(value: str):
    """Преобразует строку в Decimal."""
    try:
        return Decimal(normalize_money(value))
    except (InvalidOperation, ValueError):
        return None


def normalize_date(day: str, month: str, year: str) -> str:
    """Возвращает корректную дату в формате ДД.ММ.ГГГГ."""
    try:
        result = datetime(
            int(year),
            int(month),
            int(day)
        )

        return result.strftime("%d.%m.%Y")

    except (ValueError, TypeError):
        return ""


def create_border() -> Border:
    """Тонкая граница Excel."""
    side = Side(style="thin", color="000000")

    return Border(
        left=side,
        right=side,
        top=side,
        bottom=side
    )


# ============================================================
# РАБОТА СО СЛОВАМИ PDF
# ============================================================

def get_word_text(word: dict) -> str:
    return str(word.get("text", "")).strip()


def get_word_x0(word: dict) -> float:
    return float(word.get("x0", 0))


def get_word_x1(word: dict) -> float:
    return float(word.get("x1", 0))


def get_word_top(word: dict) -> float:
    return float(word.get("top", 0))


def get_word_bottom(word: dict) -> float:
    return float(word.get("bottom", 0))


def extract_page_words(page) -> list[dict]:
    """
    Извлекает отдельные слова и отдельные цифры.

    keep_blank_chars=False позволяет получать цифры,
    находящиеся в отдельных квадратных полях.
    """
    return page.extract_words(
        x_tolerance=1,
        y_tolerance=2,
        keep_blank_chars=False,
        use_text_flow=False
    ) or []


def group_words_into_lines(
    words: list[dict],
    tolerance: float = 4
) -> list[list[dict]]:
    """Объединяет слова в визуальные строки."""
    if not words:
        return []

    sorted_words = sorted(
        words,
        key=lambda item: (
            get_word_top(item),
            get_word_x0(item)
        )
    )

    lines: list[list[dict]] = []

    for word in sorted_words:
        word_top = get_word_top(word)

        matching_line = None

        for line in lines:
            line_top = sum(
                get_word_top(item)
                for item in line
            ) / len(line)

            if abs(word_top - line_top) <= tolerance:
                matching_line = line
                break

        if matching_line is None:
            lines.append([word])
        else:
            matching_line.append(word)

    for line in lines:
        line.sort(key=get_word_x0)

    lines.sort(
        key=lambda line: min(
            get_word_top(item)
            for item in line
        )
    )

    return lines


def join_words_preserving_lines(words: list[dict]) -> str:
    """
    Объединяет несколько визуальных строк.

    Это исправляет наименования организаций,
    которые занимают две или больше строк.
    """
    lines = group_words_into_lines(words)

    text_lines = []

    for line in lines:
        line_text = " ".join(
            get_word_text(item)
            for item in line
            if get_word_text(item)
        )

        line_text = clean_spaces(line_text)

        if line_text:
            text_lines.append(line_text)

    return clean_spaces(" ".join(text_lines))


def find_code_word(
    words: list[dict],
    code: str,
    left_half_only: bool,
    page_width: float
):
    """
    Ищет служебный код поля формы: 201, 202, 203, 206 и т.д.
    """
    candidates = []

    for word in words:
        if digits_only(get_word_text(word)) != code:
            continue

        if left_half_only and get_word_x0(word) > page_width * 0.55:
            continue

        candidates.append(word)

    if not candidates:
        return None

    return sorted(
        candidates,
        key=lambda item: (
            get_word_top(item),
            get_word_x0(item)
        )
    )[0]


def words_in_rectangle(
    words: list[dict],
    x0: float,
    x1: float,
    top: float,
    bottom: float
) -> list[dict]:
    """Возвращает слова внутри заданной области."""
    result = []

    for word in words:
        word_center_x = (
            get_word_x0(word)
            + get_word_x1(word)
        ) / 2

        word_center_y = (
            get_word_top(word)
            + get_word_bottom(word)
        ) / 2

        if (
            x0 <= word_center_x <= x1
            and top <= word_center_y <= bottom
        ):
            result.append(word)

    return result


def remove_field_labels(value: str) -> str:
    """Удаляет служебные подписи формы из результата."""
    patterns = [
        r"^\s*202\s*",
        r"^\s*Ф\.?\s*И\.?\s*О\.?\s*ИП\s*/\s*",
        r"^\s*Наименование\s+организации\s*:\s*",
        r"^\s*Ф\.?\s*И\.?\s*О\.?\s*ИП\s*/\s*"
        r"Наименование\s+организации\s*:\s*"
    ]

    for pattern in patterns:
        value = re.sub(
            pattern,
            "",
            value,
            flags=re.IGNORECASE
        )

    return clean_spaces(value)


# ============================================================
# ЧТЕНИЕ PDF
# ============================================================

def extract_pdf_data(pdf_path: Path):
    """
    Возвращает:
    - обычный текст;
    - layout-текст;
    - первую страницу;
    - слова первой страницы.
    """
    ordinary_parts = []
    layout_parts = []
    first_page = None
    first_page_words = []

    pdf = pdfplumber.open(pdf_path)

    try:
        for page_number, page in enumerate(pdf.pages):
            ordinary = page.extract_text(
                x_tolerance=2,
                y_tolerance=3
            )

            layout = page.extract_text(
                x_tolerance=2,
                y_tolerance=3,
                layout=True
            )

            if ordinary:
                ordinary_parts.append(ordinary)

            if layout:
                layout_parts.append(layout)

            if page_number == 0:
                first_page = page
                first_page_words = extract_page_words(page)

        ordinary_text = "\n".join(ordinary_parts)
        layout_text = "\n".join(layout_parts)

        # Возвращаем размеры и слова, а не объект page,
        # потому что PDF далее закрывается.
        page_info = None

        if first_page is not None:
            page_info = {
                "width": float(first_page.width),
                "height": float(first_page.height),
                "words": first_page_words
            }

        return ordinary_text, layout_text, page_info

    finally:
        pdf.close()


# ============================================================
# НОМЕР СЧЕТА-ФАКТУРЫ
# ============================================================

def extract_invoice_number(text: str) -> str:
    normalized = normalize_text(text)

    patterns = [
        r"102\s*(?:Номер|НОМЕР)\s*:\s*"
        r"([0-9]{4,}-[0-9]{3}-[0-9]{5,})",

        r"\b([0-9]{7}-[0-9]{3}-[0-9]{8})\b",

        r"\b([0-9]{4,}-[0-9]{3}-[0-9]{5,})\b"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            normalized,
            re.IGNORECASE
        )

        if match:
            return match.group(1).strip()

    return ""


# ============================================================
# ДАТА ОФОРМЛЕНИЯ
# ============================================================

def extract_issue_date(text: str) -> str:
    normalized = normalize_text(text)

    marker = re.search(
        r"103\s+Дата\s+оформления\s*:?",
        normalized,
        re.IGNORECASE
    )

    if marker:
        fragment = normalized[
            marker.end():marker.end() + 350
        ]

        normal_date = re.search(
            r"\b(\d{1,2})[./\-](\d{1,2})[./\-](\d{4})\b",
            fragment
        )

        if normal_date:
            return normalize_date(
                normal_date.group(1),
                normal_date.group(2),
                normal_date.group(3)
            )

        fragment = re.split(
            r"201\s+Поставщик",
            fragment,
            maxsplit=1,
            flags=re.IGNORECASE
        )[0]

        digits = re.findall(r"\d", fragment)

        for start in range(max(0, len(digits) - 7)):
            candidate = "".join(
                digits[start:start + 8]
            )

            result = normalize_date(
                candidate[0:2],
                candidate[2:4],
                candidate[4:8]
            )

            if result:
                return result

    return ""


# ============================================================
# ИНН ПОСТАВЩИКА
# ============================================================

def is_valid_supplier_inn(value: str) -> bool:
    """
    Основной стандарт — 14 цифр.

    Другие длины не принимаются автоматически,
    чтобы не подставить номер счёта или ИНН покупателя.
    """
    inn = digits_only(value)

    return (
        len(inn) == 14
        and inn != "00000000000000"
    )


def extract_supplier_inn_from_coordinates(
    page_info
) -> tuple[str, str]:
    """
    Извлекает ИНН из строки 201 по координатам.

    Поддерживает:
    - 03004202410293 одним текстом;
    - 0 3 0 0 4 2 ... отдельными цифрами;
    - цифры в отдельных квадратных полях.
    """
    if not page_info:
        return "", ""

    words = page_info["words"]
    width = page_info["width"]

    code_201 = find_code_word(
        words,
        "201",
        left_half_only=True,
        page_width=width
    )

    code_202 = find_code_word(
        words,
        "202",
        left_half_only=True,
        page_width=width
    )

    if not code_201:
        return "", ""

    row_top = get_word_top(code_201) - 5

    if code_202:
        row_bottom = get_word_top(code_202) - 1
    else:
        row_bottom = get_word_bottom(code_201) + 25

    row_words = words_in_rectangle(
        words,
        x0=get_word_x1(code_201),
        x1=width * 0.55,
        top=row_top,
        bottom=row_bottom
    )

    # На случай обычного цельного ИНН.
    row_text = " ".join(
        get_word_text(word)
        for word in sorted(
            row_words,
            key=get_word_x0
        )
    )

    normal_candidates = re.findall(
        r"(?<!\d)\d{14}(?!\d)",
        row_text
    )

    for candidate in normal_candidates:
        if is_valid_supplier_inn(candidate):
            return candidate, "ИНН найден в поле 201 по координатам"

    # Для цифр, расположенных по одной в ячейках.
    digit_tokens = []

    for word in sorted(row_words, key=get_word_x0):
        token = digits_only(get_word_text(word))

        if not token:
            continue

        # Исключаем сам номер поля и случайные длинные значения.
        if token == "201":
            continue

        if len(token) == 1:
            digit_tokens.append(token)

        elif len(token) == 14:
            if is_valid_supplier_inn(token):
                return token, "ИНН найден целым значением в поле 201"

        elif 2 <= len(token) <= 13:
            # Иногда pdfplumber объединяет несколько соседних ячеек.
            digit_tokens.extend(list(token))

    # Метка "Поставщик ИНН" может содержать цифр не должна.
    # Берём любую последовательность из 14 цифр.
    joined_digits = "".join(digit_tokens)

    for start in range(max(0, len(joined_digits) - 13)):
        candidate = joined_digits[start:start + 14]

        if is_valid_supplier_inn(candidate):
            return candidate, "ИНН собран из отдельных квадратных полей"

    return "", ""


def extract_supplier_inn_from_text(
    ordinary_text: str,
    layout_text: str
) -> tuple[str, str]:
    """Резервное извлечение ИНН из текста."""
    combined = normalize_text(
        ordinary_text + "\n" + layout_text
    )

    patterns = [
        r"201\s+Поставщик\s+ИНН\s*:\s*(\d{14})",

        r"201\s+Поставщик\s+ИНН\s*:\s*"
        r"((?:\d[\s|]*){14})",

        r"\bПоставщик\s+ИНН\s*:?\s*"
        r"((?:\d[\s|]*){14})"
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            combined,
            re.IGNORECASE
        )

        if not match:
            continue

        candidate = digits_only(match.group(1))

        if is_valid_supplier_inn(candidate):
            return candidate, "ИНН найден по тексту поля 201"

    # Ищем участок поставщика до поля покупателя.
    supplier_marker = re.search(
        r"(?:201\s*)?Поставщик\s+ИНН",
        combined,
        re.IGNORECASE
    )

    if supplier_marker:
        fragment = combined[
            supplier_marker.end():
            supplier_marker.end() + 350
        ]

        fragment = re.split(
            r"301\s+Покупатель\s+ИНН",
            fragment,
            maxsplit=1,
            flags=re.IGNORECASE
        )[0]

        # Цельные 14 цифр.
        candidates = re.findall(
            r"(?<!\d)\d{14}(?!\d)",
            fragment
        )

        for candidate in candidates:
            if is_valid_supplier_inn(candidate):
                return candidate, "ИНН найден в разделе поставщика"

        # Раздельные цифры.
        digit_tokens = re.findall(r"\d", fragment)

        for start in range(max(0, len(digit_tokens) - 13)):
            candidate = "".join(
                digit_tokens[start:start + 14]
            )

            if is_valid_supplier_inn(candidate):
                return candidate, "ИНН собран из раздельных цифр текста"

    return "", "ИНН поставщика не найден"


def extract_supplier_inn(
    ordinary_text: str,
    layout_text: str,
    page_info
) -> tuple[str, str]:
    """
    Координатный поиск имеет приоритет,
    потому что он не смешивает поставщика и покупателя.
    """
    inn, status = extract_supplier_inn_from_coordinates(
        page_info
    )

    if inn:
        return inn, status

    return extract_supplier_inn_from_text(
        ordinary_text,
        layout_text
    )


# ============================================================
# НАИМЕНОВАНИЕ ПОСТАВЩИКА
# ============================================================

def simplify_organization_name(value: str) -> str:
    value = clean_spaces(value)

    replacements = [
        (
            r"^Общество\s+с\s+ограниченной\s+"
            r"ответственностью\s*",
            "ООО "
        ),
        (
            r"^Закрытое\s+акционерное\s+общество\s*",
            "ЗАО "
        ),
        (
            r"^Открытое\s+акционерное\s+общество\s*",
            "ОАО "
        ),
        (
            r"^Индивидуальный\s+предприниматель\s*",
            "ИП "
        )
    ]

    for pattern, replacement in replacements:
        value = re.sub(
            pattern,
            replacement,
            value,
            flags=re.IGNORECASE
        )

    return clean_spaces(value)


# ============================================================
# НАИМЕНОВАНИЕ ПОСТАВЩИКА
# ============================================================

def simplify_organization_name(value: str) -> str:
    """
    Сокращает организационно-правовую форму,
    сохраняя полное название организации.
    """
    value = clean_spaces(value)

    replacements = [
        (
            r"^Общество\s+с\s+ограниченной\s+ответственностью\s*",
            "ООО "
        ),
        (
            r"^Закрытое\s+акционерное\s+общество\s*",
            "ЗАО "
        ),
        (
            r"^Открытое\s+акционерное\s+общество\s*",
            "ОАО "
        ),
        (
            r"^Муниципальное\s+предприятие\s*",
            "МП "
        ),
        (
            r"^Государственное\s+предприятие\s*",
            "ГП "
        ),
        (
            r"^Индивидуальный\s+предприниматель\s*",
            "ИП "
        )
    ]

    for pattern, replacement in replacements:
        value = re.sub(
            pattern,
            replacement,
            value,
            flags=re.IGNORECASE
        )

    return clean_spaces(value)


def find_supplier_column_right_edge(
    words: list[dict],
    page_width: float
) -> float:
    """
    Определяет вертикальную границу между поставщиком
    и покупателем.

    Ищет коды правой колонки: 301, 302, 303, 304.
    Это надёжнее, чем использовать половину ширины страницы.
    """
    possible_edges = []

    for word in words:
        text = digits_only(get_word_text(word))

        if text not in {"301", "302", "303", "304", "305", "306"}:
            continue

        word_x0 = get_word_x0(word)

        # Правая колонка обычно начинается после 45% страницы.
        if word_x0 > page_width * 0.44:
            possible_edges.append(word_x0)

    if possible_edges:
        return min(possible_edges) - 2

    # Резервная граница.
    return page_width * 0.52


def extract_supplier_name_from_coordinates(
    page_info
) -> str:
    """
    Извлекает название поставщика, включая продолжение
    на второй, третьей и последующих строках.

    Чтение начинается со строки:
    Ф.И.О. ИП/Наименование организации:

    и заканчивается перед строкой:
    Филиал поставщика ИНН
    """
    if not page_info:
        return ""

    words = page_info["words"]
    page_width = page_info["width"]

    supplier_right_edge = find_supplier_column_right_edge(
        words,
        page_width
    )

    # Оставляем только левую колонку поставщика.
    supplier_words = []

    for word in words:
        center_x = (
            get_word_x0(word)
            + get_word_x1(word)
        ) / 2

        if center_x < supplier_right_edge:
            supplier_words.append(word)

    lines = group_words_into_lines(
        supplier_words,
        tolerance=3.5
    )

    collecting = False
    name_parts = []

    for line in lines:
        line_words = sorted(
            line,
            key=get_word_x0
        )

        line_text = " ".join(
            get_word_text(word)
            for word in line_words
            if get_word_text(word)
        )

        line_text = clean_spaces(line_text)
        lower_text = line_text.lower()

        if not line_text:
            continue

        if not collecting:
            # Начало поля наименования организации.
            is_name_field = (
                "наименование" in lower_text
                and "организац" in lower_text
                and (
                    "ф.и.о" in lower_text
                    or "фио" in lower_text
                    or "ип/" in lower_text
                )
            )

            if not is_name_field:
                continue

            collecting = True

            # Удаляем номер поля 202, в том числе вариант:
            # 2 0 2
            line_text = re.sub(
                r"^\s*2\s*0\s*2\s*",
                "",
                line_text
            )

            # Берём всё после подписи
            # «Наименование организации:».
            match = re.search(
                r"Наименование\s+организац(?:ии|ия)"
                r"\s*:?\s*(.*)$",
                line_text,
                re.IGNORECASE
            )

            if match:
                first_part = clean_spaces(
                    match.group(1)
                )

                if first_part:
                    name_parts.append(first_part)

            continue

        # Конец поля поставщика.
        stop_markers = [
            "филиал поставщика",
            "наименование филиала",
            "адрес (юридич",
            "код и наименование налогового органа"
        ]

        if any(
            marker in lower_text
            for marker in stop_markers
        ):
            break

        # Также останавливаемся на кодах следующих полей.
        compact_digits = digits_only(line_text)

        if compact_digits in {
            "203",
            "204",
            "205",
            "206"
        }:
            break

        # Убираем случайно попавший номер поля.
        continuation = re.sub(
            r"^\s*2\s*0\s*[3-9]\s*",
            "",
            line_text
        )

        continuation = clean_spaces(continuation)

        # Убираем код правой колонки, если он попал
        # в конец строки.
        continuation = re.sub(
            r"\s+30[1-9]\s*$",
            "",
            continuation
        )

        continuation = clean_spaces(continuation)

        if continuation:
            name_parts.append(continuation)

    result = clean_spaces(
        " ".join(name_parts)
    )

    # Дополнительная очистка служебных кодов.
    result = re.sub(
        r"^\s*202\s*",
        "",
        result
    )

    result = re.sub(
        r"\s+30[1-9]\s*$",
        "",
        result
    )

    result = clean_spaces(result)

    if not result:
        return ""

    return simplify_organization_name(result)


def extract_supplier_name_from_text(
    text: str
) -> str:
    """
    Резервный способ для PDF, где координаты слов
    извлекаются некорректно.

    Поддерживает перенос названия на несколько строк.
    """
    if not text:
        return ""

    lines = text.splitlines()

    collecting = False
    result_parts = []

    for original_line in lines:
        line = clean_spaces(original_line)
        lower_line = line.lower()

        if not line:
            continue

        if not collecting:
            if (
                "наименование" in lower_line
                and "организац" in lower_line
                and (
                    "ф.и.о" in lower_line
                    or "фио" in lower_line
                    or "ип/" in lower_line
                )
            ):
                collecting = True

                # Левая часть до поля покупателя.
                line = re.split(
                    r"\b302\b",
                    line,
                    maxsplit=1
                )[0]

                match = re.search(
                    r"Наименование\s+организац(?:ии|ия)"
                    r"\s*:?\s*(.*)$",
                    line,
                    re.IGNORECASE
                )

                if match:
                    part = clean_spaces(
                        match.group(1)
                    )

                    if part:
                        result_parts.append(part)

                continue

        else:
            # Прекращаем чтение возле следующего поля.
            if (
                "филиал поставщика" in lower_line
                or "наименование филиала" in lower_line
                or re.search(
                    r"^\s*2\s*0\s*3\b",
                    line
                )
                or re.search(
                    r"^\s*203\b",
                    line
                )
            ):
                break

            # Не захватываем правую колонку покупателя.
            line = re.split(
                r"\b302\b",
                line,
                maxsplit=1
            )[0]

            line = re.sub(
                r"\s+30[1-9]\s*$",
                "",
                line
            )

            part = clean_spaces(line)

            if part:
                result_parts.append(part)

    result = clean_spaces(
        " ".join(result_parts)
    )

    if not result:
        return ""

    return simplify_organization_name(result)


def extract_supplier_name(
    ordinary_text: str,
    layout_text: str,
    page_info
) -> str:
    """
    Последовательно использует координаты,
    обычный текст и layout-текст.
    """
    result = extract_supplier_name_from_coordinates(
        page_info
    )

    if result:
        return result

    result = extract_supplier_name_from_text(
        ordinary_text
    )

    if result:
        return result

    return extract_supplier_name_from_text(
        layout_text
    )

# ============================================================
# КОД НАЛОГОВОГО ОРГАНА
# ============================================================

def extract_tax_code_from_coordinates(page_info) -> str:
    """
    Извлекает первые три цифры из строки 206 поставщика.

    Работает и для:
    004 - УГНС...
    и для отдельных клеток:
    9 9 9 УККН...
    """
    if not page_info:
        return ""

    words = page_info["words"]
    width = page_info["width"]

    code_206 = find_code_word(
        words,
        "206",
        left_half_only=True,
        page_width=width
    )

    code_207 = find_code_word(
        words,
        "207",
        left_half_only=True,
        page_width=width
    )

    if not code_206:
        return ""

    top = get_word_top(code_206) - 4

    if code_207:
        bottom = get_word_top(code_207) - 1
    else:
        bottom = get_word_bottom(code_206) + 30

    row_words = words_in_rectangle(
        words,
        x0=get_word_x1(code_206),
        x1=width * 0.55,
        top=top,
        bottom=bottom
    )

    # Сначала цельный код 004 или 999.
    text = " ".join(
        get_word_text(word)
        for word in sorted(row_words, key=get_word_x0)
    )

    normal_match = re.search(
        r"(?<!\d)(\d{3})(?!\d)",
        text
    )

    if normal_match:
        return normal_match.group(1)

    # Затем отдельные цифры в ячейках.
    digits = []

    for word in sorted(row_words, key=get_word_x0):
        token = digits_only(get_word_text(word))

        if len(token) == 1:
            digits.append(token)

        elif len(token) == 3:
            return token

    if len(digits) >= 3:
        return "".join(digits[:3])

    return ""


def extract_tax_office_code(
    ordinary_text: str,
    layout_text: str,
    page_info
) -> str:
    code = extract_tax_code_from_coordinates(
        page_info
    )

    if code:
        return code.zfill(3)

    combined = normalize_text(
        ordinary_text + "\n" + layout_text
    )

    patterns = [
        (
            r"Код\s+и\s+наименование\s+налогового\s+органа\s*:"
            r"\s*(\d{3})\s*-"
        ),
        (
            r"206.*?налогового\s+органа\s*:"
            r"\s*((?:\d\s*){3})"
        ),
        (
            r"(?<!\d)(\d\s+\d\s+\d)"
            r"\s+(?:УГНС|УККН)"
        )
    ]

    for pattern in patterns:
        match = re.search(
            pattern,
            combined,
            re.IGNORECASE | re.DOTALL
        )

        if match:
            candidate = digits_only(match.group(1))

            if len(candidate) == 3:
                return candidate

    return ""


# ============================================================
# СУММА НДС
# ============================================================

def extract_vat_from_total(text: str):
    """
    Поддерживает оба варианта:

    41777.35 5013.28 2088.87 48879.50

    3183.37 X 382.00 X 31.83 3597.20 X
    """
    normalized = normalize_text(text)

    markers = [
        r"Итого\s+по\s+счету\s*-\s*фактуре\s*:",
        r"Итого\s+по\s+счету-фактуре\s*:",
        r"Итого\s+по\s+счетуфактуре\s*:"
    ]

    marker = None

    for pattern in markers:
        marker = re.search(
            pattern,
            normalized,
            re.IGNORECASE
        )

        if marker:
            break

    if not marker:
        return None

    fragment = normalized[
        marker.end():marker.end() + 300
    ]

    values = re.findall(
        r"(?<!\d)"
        r"\d[\d \u00a0]*[.,]\d{2,5}"
        r"(?!\d)",
        fragment
    )

    parsed = []

    for value in values:
        number = parse_decimal(value)

        if number is not None:
            parsed.append(number)

    # Первое число — сумма без НДС.
    # Второе число — сумма НДС.
    if len(parsed) >= 2:
        return parsed[1]

    return None


# ============================================================
# НОМЕР ПЛАТЕЖНОГО ВАУЧЕРА ИЗ ИМЕНИ PDF
# ============================================================

def extract_payment_voucher(pdf_path: Path) -> str:
    """
    Извлекает номер платежного ваучера из имени PDF.

    Формат: 53112605-INV.pdf -> 53112605.
    Если имя не соответствует этому формату, возвращается пусто,
    чтобы случайно не принять другой номер за ваучер.
    """
    stem = pdf_path.stem.strip()
    match = re.match(r"^(\d+)-INV$", stem, re.IGNORECASE)
    return match.group(1) if match else ""


# ============================================================
# ОБРАБОТКА ОДНОГО PDF
# ============================================================

def process_pdf(pdf_path: Path) -> dict:
    ordinary_text, layout_text, page_info = extract_pdf_data(
        pdf_path
    )

    if not ordinary_text.strip() and not layout_text.strip():
        raise ValueError(
            "PDF не содержит извлекаемого текста"
        )

    combined_text = ordinary_text + "\n" + layout_text

    supplier_inn, inn_status = extract_supplier_inn(
        ordinary_text,
        layout_text,
        page_info
    )

    supplier = extract_supplier_name(
        ordinary_text,
        layout_text,
        page_info
    )

    tax_code = extract_tax_office_code(
        ordinary_text,
        layout_text,
        page_info
    )

    invoice_number = extract_invoice_number(
        combined_text
    )

    issue_date = extract_issue_date(
        combined_text
    )

    vat = extract_vat_from_total(
        combined_text
    )

    missing = []

    if not supplier:
        missing.append("наименование поставщика")

    if not supplier_inn:
        missing.append("ИНН поставщика")

    if not tax_code:
        missing.append("код налогового органа")

    if not invoice_number:
        missing.append("номер СФ")

    if not issue_date:
        missing.append("дата")

    if vat is None:
        missing.append("сумма НДС")

    if missing:
        status = (
            "Не найдено: "
            + ", ".join(missing)
            + ". "
            + inn_status
        )
    else:
        status = (
            "Все поля найдены. "
            + inn_status
        )

    return {
        "source_file": pdf_path.name,
        "payment_voucher": extract_payment_voucher(pdf_path),
        "supplier": supplier,
        "supplier_inn": supplier_inn,
        "tax_office_code": tax_code,
        "invoice_number": invoice_number,
        "issue_date": issue_date,
        "vat": vat,
        "status": status,
        "is_complete": len(missing) == 0
    }


# ============================================================
# СОЗДАНИЕ EXCEL
# ============================================================

def save_excel(
    rows: list[dict],
    errors: list[dict],
    output_path: Path
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "НДС"

    border = create_border()

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    warning_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )

    error_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC"
    )

    # Заголовок VAAT
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("B1:B2")
    sheet.merge_cells("C1:C2")
    sheet.merge_cells("D1:D2")
    sheet.merge_cells("E1:G1")
    sheet.merge_cells("H1:H3")

    sheet["A1"] = "№"
    sheet["B1"] = (
        "Наименование поставщика "
        "товаров (работ, услуг)"
    )
    sheet["C1"] = "ИНН"
    sheet["D1"] = "Код налогового органа"
    sheet["E1"] = "Счет-фактура"
    sheet["H1"] = "Payment Voucher#"

    sheet["E2"] = "№"
    sheet["F2"] = "дата"
    sheet["G2"] = "сумма НДС (сом.)"

    for column in range(1, 9):
        sheet.cell(
            row=3,
            column=column,
            value=column
        )

    for row_cells in sheet.iter_rows(
        min_row=1,
        max_row=3,
        min_col=1,
        max_col=7
    ):
        for cell in row_cells:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

    # Данные
    for index, item in enumerate(rows, start=1):
        sheet.append([
            index,
            item["supplier"],
            item["supplier_inn"],
            item["tax_office_code"],
            item["invoice_number"],
            item["issue_date"],
            (
                float(item["vat"])
                if item["vat"] is not None
                else None
            ),
            item["payment_voucher"]
        ])

        excel_row = sheet.max_row

        for column in range(1, 9):
            cell = sheet.cell(
                row=excel_row,
                column=column
            )

            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

            if not item["is_complete"]:
                cell.fill = warning_fill

        # Текстовый формат сохраняет ведущие нули.
        for column in (3, 4, 5, 6, 8):
            sheet.cell(
                row=excel_row,
                column=column
            ).number_format = "@"

        sheet.cell(
            row=excel_row,
            column=7
        ).number_format = '#,##0.00'

        if not item["supplier_inn"]:
            inn_cell = sheet.cell(
                row=excel_row,
                column=3
            )

            inn_cell.value = "НЕ НАЙДЕН"
            inn_cell.fill = error_fill

    # Итог
    total_row = sheet.max_row + 1

    sheet.merge_cells(
        start_row=total_row,
        start_column=1,
        end_row=total_row,
        end_column=6
    )

    sheet.cell(
        row=total_row,
        column=1,
        value="Итого"
    )

    if total_row > 4:
        sheet.cell(
            row=total_row,
            column=7,
            value=f"=SUM(G4:G{total_row - 1})"
        )
    else:
        sheet.cell(
            row=total_row,
            column=7,
            value=0
        )

    for column in range(1, 9):
        cell = sheet.cell(
            row=total_row,
            column=column
        )

        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    sheet.cell(
        row=total_row,
        column=7
    ).number_format = '#,##0.00'

    widths = {
        "A": 7,
        "B": 50,
        "C": 20,
        "D": 23,
        "E": 29,
        "F": 16,
        "G": 20,
        "H": 20
    }

    for column_letter, width in widths.items():
        sheet.column_dimensions[
            column_letter
        ].width = width

    sheet.row_dimensions[1].height = 40
    sheet.row_dimensions[2].height = 27
    sheet.row_dimensions[3].height = 22

    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:H{total_row - 1}"
    sheet.sheet_view.showGridLines = False

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    # Лист проверки
    check_sheet = workbook.create_sheet("Проверка")

    check_headers = [
        "Файл PDF",
        "Payment Voucher#",
        "Наименование поставщика",
        "ИНН",
        "Код налогового органа",
        "Номер СФ",
        "Дата",
        "Сумма НДС",
        "Результат проверки"
    ]

    check_sheet.append(check_headers)

    for cell in check_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for item in rows:
        check_sheet.append([
            item["source_file"],
            item["payment_voucher"],
            item["supplier"],
            item["supplier_inn"],
            item["tax_office_code"],
            item["invoice_number"],
            item["issue_date"],
            (
                float(item["vat"])
                if item["vat"] is not None
                else None
            ),
            item["status"]
        ])

        current_row = check_sheet.max_row

        for column in range(1, 10):
            cell = check_sheet.cell(
                row=current_row,
                column=column
            )

            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

            if not item["is_complete"]:
                cell.fill = warning_fill

        for column in (2, 4, 5, 6, 7):
            check_sheet.cell(
                row=current_row,
                column=column
            ).number_format = "@"

        check_sheet.cell(
            row=current_row,
            column=8
        ).number_format = '#,##0.00'

    check_widths = {
        "A": 43,
        "B": 20,
        "C": 50,
        "D": 20,
        "E": 23,
        "F": 29,
        "G": 16,
        "H": 18,
        "I": 65
    }

    for column_letter, width in check_widths.items():
        check_sheet.column_dimensions[
            column_letter
        ].width = width

    check_sheet.freeze_panes = "A2"
    check_sheet.auto_filter.ref = (
        f"A1:I{check_sheet.max_row}"
    )

    # Ошибки открытия
    if errors:
        error_sheet = workbook.create_sheet("Ошибки")
        error_sheet.append(["Файл PDF", "Ошибка"])

        for cell in error_sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = error_fill
            cell.border = border

        for error in errors:
            error_sheet.append([
                error["file"],
                error["error"]
            ])

        error_sheet.column_dimensions["A"].width = 45
        error_sheet.column_dimensions["B"].width = 90

    workbook.save(output_path)


# ============================================================
# ОКНО ПРОГРЕССА
# ============================================================

class ProgressWindow:
    def __init__(self, root: Tk, total_files: int):
        self.root = root
        self.total_files = total_files
        self.cancelled = BooleanVar(value=False)

        self.window = Toplevel(root)
        self.window.title("Извлечение данных из PDF")
        self.window.geometry("560x210")
        self.window.resizable(False, False)
        self.window.protocol(
            "WM_DELETE_WINDOW",
            self.request_cancel
        )

        self.window.attributes("-topmost", True)

        self.title_label = Label(
            self.window,
            text="Обработка счетов-фактур",
            font=("Segoe UI", 14, "bold")
        )
        self.title_label.pack(pady=(18, 8))

        self.file_text = StringVar(
            value="Подготовка к обработке..."
        )

        self.file_label = Label(
            self.window,
            textvariable=self.file_text,
            font=("Segoe UI", 9),
            wraplength=510,
            justify="center"
        )
        self.file_label.pack(pady=(0, 12))

        self.progress = ttk.Progressbar(
            self.window,
            orient="horizontal",
            mode="determinate",
            maximum=max(total_files, 1),
            length=500
        )
        self.progress.pack(pady=4)

        self.counter_text = StringVar(
            value=f"0 из {total_files}"
        )

        self.counter_label = Label(
            self.window,
            textvariable=self.counter_text,
            font=("Segoe UI", 10)
        )
        self.counter_label.pack(pady=7)

        self.cancel_button = Button(
            self.window,
            text="Отменить",
            width=15,
            command=self.request_cancel
        )
        self.cancel_button.pack(pady=4)

        self.window.update_idletasks()

    def request_cancel(self):
        self.cancelled.set(True)
        self.file_text.set(
            "Отмена после завершения текущего файла..."
        )
        self.cancel_button.config(state="disabled")
        self.window.update_idletasks()

    def update(
        self,
        current: int,
        filename: str
    ):
        self.progress["value"] = current
        self.counter_text.set(
            f"{current} из {self.total_files}"
        )
        self.file_text.set(
            f"Обрабатывается:\n{filename}"
        )

        self.window.update_idletasks()
        self.window.update()

    def set_saving(self):
        self.file_text.set(
            "Формируется файл Result.xlsx..."
        )
        self.window.update_idletasks()
        self.window.update()

    def close(self):
        try:
            self.window.destroy()
        except Exception:
            pass


# ============================================================
# ИНТЕРФЕЙС
# ============================================================

def choose_folder(root: Tk, default_folder: Path):
    selected = filedialog.askdirectory(
        parent=root,
        title="Выберите папку с PDF-счетами-фактурами",
        initialdir=str(default_folder)
    )

    if not selected:
        return None

    return Path(selected)


def show_message(
    root: Tk,
    title: str,
    text: str,
    error: bool = False
):
    if error:
        messagebox.showerror(
            title,
            text,
            parent=root
        )
    else:
        messagebox.showinfo(
            title,
            text,
            parent=root
        )


# ============================================================
# ЗАПУСК
# ============================================================

def main():
    root = Tk()
    root.withdraw()

    selected_folder = choose_folder(
        root,
        application_directory()
    )

    if selected_folder is None:
        root.destroy()
        return

    # Поддерживаем .pdf и .PDF.
    pdf_files = sorted(
        [
            path
            for path in selected_folder.iterdir()
            if (
                path.is_file()
                and path.suffix.lower() == ".pdf"
            )
        ],
        key=lambda path: path.name.lower()
    )

    if not pdf_files:
        show_message(
            root,
            "PDF не найдены",
            "В выбранной папке нет PDF-файлов.",
            error=True
        )
        root.destroy()
        return

    progress_window = ProgressWindow(
        root,
        len(pdf_files)
    )

    rows = []
    errors = []
    processed_count = 0

    for index, pdf_path in enumerate(
        pdf_files,
        start=1
    ):
        if progress_window.cancelled.get():
            break

        progress_window.update(
            index - 1,
            pdf_path.name
        )

        try:
            rows.append(
                process_pdf(pdf_path)
            )

        except Exception as exc:
            errors.append({
                "file": pdf_path.name,
                "error": str(exc)
            })

        processed_count += 1

        progress_window.update(
            index,
            pdf_path.name
        )

    if processed_count == 0:
        progress_window.close()
        root.destroy()
        return

    progress_window.set_saving()

    output_path = selected_folder / "Result.xlsx"

    try:
        save_excel(
            rows,
            errors,
            output_path
        )

    except PermissionError:
        progress_window.close()

        show_message(
            root,
            "Не удалось сохранить Result.xlsx",
            "Файл Result.xlsx открыт в Excel.\n\n"
            "Закройте его и запустите программу повторно.",
            error=True
        )

        root.destroy()
        return

    except Exception as exc:
        progress_window.close()

        show_message(
            root,
            "Ошибка сохранения",
            f"Не удалось создать Result.xlsx:\n\n{exc}",
            error=True
        )

        root.destroy()
        return

    progress_window.close()

    complete_count = sum(
        1
        for row in rows
        if row["is_complete"]
    )

    missing_inn_count = sum(
        1
        for row in rows
        if not row["supplier_inn"]
    )

    cancelled_text = ""

    if processed_count < len(pdf_files):
        cancelled_text = (
            "\n\nОбработка была отменена пользователем."
        )

    result_message = (
        f"Обработано PDF: {processed_count}\n"
        f"Все поля найдены: {complete_count}\n"
        f"ИНН поставщика не найден: "
        f"{missing_inn_count}\n"
        f"Ошибок открытия: {len(errors)}\n\n"
        f"Результат:\n{output_path}"
        f"{cancelled_text}"
    )

    show_message(
        root,
        "Обработка завершена",
        result_message
    )

    root.destroy()


if __name__ == "__main__":
    try:
        main()

    except Exception:
        error_path = (
            application_directory()
            / "Extract_error.txt"
        )

        error_path.write_text(
            traceback.format_exc(),
            encoding="utf-8"
        )

        try:
            error_root = Tk()
            error_root.withdraw()

            messagebox.showerror(
                "Ошибка программы",
                "Произошла непредвиденная ошибка.\n\n"
                "Подробности записаны в файл:\n"
                f"{error_path}",
                parent=error_root
            )

            error_root.destroy()

        except Exception:
            pass
